"""Topical quiz regressions. RUN_RELEASE_TESTS=1 and disposable @auditdb required."""
import asyncio
import datetime as dt
import io
import os
import unittest

import asyncpg
from docx import Document
from fastapi import HTTPException
from core.config import DATABASE_URL
from core.quiz_attempts import finalize_expired_quizzes
from routers.quizzes import parse_docx_questions
from scripts.init_db import SCHEMA, ENT_MIGRATIONS, RELEASE_MIGRATIONS
from scripts.topical_migrations import TOPICAL_MIGRATIONS
from tests import test_release as release


def document(lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class DraftWordTests(unittest.TestCase):
    def test_missing_markers_are_drafts_not_first_answer(self):
        questions = parse_docx_questions(document(['<question>', 'Қазақстанның астанасы?', '<variant>Алматы', '<variant>Астана']))
        self.assertEqual(len(questions), 1)
        self.assertFalse(questions[0]['is_complete'])
        self.assertFalse(any(option['is_correct'] for option in questions[0]['options']))

    def test_mixed_and_multiple_correct_markers(self):
        questions = parse_docx_questions(document(['<question>A', '<variant_correct>1', '<variant>2', '<question>B', '<variant>3', '<variant>4', '<question>C', '<variant_correct>5', '<variant_correct>6']))
        self.assertEqual([q['is_complete'] for q in questions], [True, False, True])
        self.assertEqual(questions[2]['type'], 'multiple')

    def test_malformed_markup_and_empty_options_remain_errors(self):
        for lines in [['<question>', '<variant>1', '<variant>2'], ['<question>A', '<variant>', '<variant>2'], ['<question>A', '<unsupported>x', '<variant>1', '<variant>2'], ['<variant>orphan', '<question>A', '<variant>1', '<variant>2']]:
            with self.subTest(lines=lines), self.assertRaises(HTTPException):
                parse_docx_questions(document(lines))

    def test_interleaved_table_keeps_question_order(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = '<question>Table\n<variant>1\n<variant>2'
        for line in ['<question>Paragraph', '<variant_correct>3', '<variant>4']:
            doc.add_paragraph(line)
        output = io.BytesIO(); doc.save(output)
        self.assertEqual([q['prompt'] for q in parse_docx_questions(output.getvalue())], ['Table', 'Paragraph'])


@unittest.skipUnless(os.environ.get('RUN_RELEASE_TESTS') == '1', 'Disposable auditdb required')
class TopicalTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = release.ReleaseTests.asyncSetUp
    cleanup = release.ReleaseTests.cleanup
    request = release.ReleaseTests.request
    item = release.ReleaseTests.item
    quiz = release.ReleaseTests.quiz

    async def new_quiz(self, attempts=1, deadline=None):
        response = await self.request('POST', f'/courses/{self.course}/items', 'teacher', json={'type': 'quiz', 'title': 'Topical fixture'})
        self.assertEqual(response.status_code, 200, response.text)
        iid = response.json()['id']
        permission = await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json={
            'student_ids': [self.users['student']['id']], 'quiz_rules': {'max_attempts': attempts, 'deadline_at': deadline}})
        self.assertEqual(permission.status_code, 200, permission.text)
        return iid

    async def question(self, iid, points=1):
        response = await self.request('POST', f'/quizzes/{iid}/questions', 'teacher', json={
            'type': 'single', 'prompt': 'Question', 'points': points, 'options': [{'text': 'Right', 'is_correct': True}, {'text': 'Wrong'}]})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['id']

    async def submit_correct(self, iid, number=1):
        quiz = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()
        answers = [{'question_id': q['id'], 'selected_option_ids': [o['id'] for o in q['options'] if o['is_correct']]} for q in quiz['questions']]
        return await self.request('POST', f'/quizzes/{iid}/attempt', json={'attempt_number': number, 'revision': 0, 'answers': answers})

    async def test_draft_import_completion_readiness_and_grading(self):
        iid = await self.new_quiz()
        content = document(['<question>Capital?', '<variant>Almaty', '<variant>Astana'])
        files = {'file': ('draft.docx', content)}
        preview = (await self.request('POST', f'/quizzes/{iid}/import-docx?preview=true', 'teacher', files=files)).json()
        self.assertEqual((preview['count'], preview['with_answers'], preview['missing_answers'], preview['errors']), (1, 0, 1, []))
        imported = await self.request('POST', f'/quizzes/{iid}/import-docx?preview=false', 'teacher', files=files)
        self.assertEqual(imported.status_code, 200, imported.text)
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 409)
        draft = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()
        self.assertFalse(draft['ready'])
        question = draft['questions'][0]
        self.assertFalse(question['is_complete'])
        option = next(o for o in question['options'] if o['text'] == 'Astana')
        denied = await self.request('PATCH', f'/quizzes/questions/{question["id"]}/correct-answers', json={'type': 'single', 'option_ids': [option['id']]})
        self.assertEqual(denied.status_code, 403)
        updated = await self.request('PATCH', f'/quizzes/questions/{question["id"]}/correct-answers', 'teacher', json={'type': 'single', 'option_ids': [option['id']]})
        self.assertEqual(updated.status_code, 200, updated.text)
        refreshed = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()
        self.assertTrue(refreshed['ready'])
        self.assertTrue(next(o for o in refreshed['questions'][0]['options'] if o['id'] == option['id'])['is_correct'])
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 200)
        submitted = await self.submit_correct(iid)
        self.assertEqual(submitted.json()['auto_score'], 1)
        frozen = await self.request('PATCH', f'/quizzes/questions/{question["id"]}/correct-answers', 'teacher', json={'type': 'single', 'option_ids': []})
        self.assertEqual(frozen.status_code, 409)

    async def test_import_deduplicates_concurrent_and_repeated_requests(self):
        iid = await self.new_quiz()
        content = document(['<question>A', '<variant_correct>1', '<variant>2'])
        async def import_once():
            return await self.request('POST', f'/quizzes/{iid}/import-docx?preview=false', 'teacher', files={'file': ('same.docx', content)})
        responses = await asyncio.gather(import_once(), import_once())
        self.assertTrue(all(r.status_code == 200 for r in responses), [r.text for r in responses])
        self.assertEqual(sum(r.json()['imported'] for r in responses), 1)
        self.assertTrue((await import_once()).json()['duplicate'])
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM quiz_questions WHERE quiz_id=$1', iid), 1)

    async def test_malformed_preview_reports_errors_and_import_is_atomic(self):
        iid = await self.new_quiz()
        files = {'file': ('broken.docx', document(['<question>A', '<variant>']))}
        preview = (await self.request('POST', f'/quizzes/{iid}/import-docx?preview=true', 'teacher', files=files)).json()
        self.assertTrue(preview['errors'])
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/import-docx?preview=false', 'teacher', files=files)).status_code, 422)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM quiz_questions WHERE quiz_id=$1', iid), 0)

    async def test_points_maximum_settings_and_legacy_lesson_settings(self):
        iid = await self.new_quiz()
        await self.question(iid, 3); await self.question(iid, 5)
        setting = await self.request('PUT', f'/quizzes/{iid}/settings', 'teacher', json={'max_score': 999, 'weight_pct': 99, 'time_limit_minutes': 4, 'shuffle_questions': False, 'shuffle_options': False})
        self.assertEqual(setting.status_code, 200, setting.text)
        view = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()
        self.assertEqual((view['max_score'], view['weight_pct'], view['time_limit_minutes']), (8, 0, 4))
        self.assertFalse(view['shuffle_questions']); self.assertFalse(view['shuffle_options'])
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 200)
        result = await self.submit_correct(iid)
        self.assertEqual(result.json()['auto_score'], 8)
        lesson = await self.item('lesson')
        response = await self.request('PUT', f'/quizzes/{lesson}/settings', 'teacher', json={'max_score': 50, 'weight_pct': 25})
        self.assertEqual(response.status_code, 200, response.text)
        row = await self.pool.fetchrow('SELECT * FROM quizzes WHERE item_id=$1', lesson)
        self.assertEqual((row['max_score'], float(row['weight_pct'])), (50, 25))

    async def test_limits_one_two_three_and_concurrent_start_and_stale_tabs(self):
        for limit in (1, 2, 3):
            iid = await self.new_quiz(limit)
            await self.question(iid, 2)
            for number in range(1, limit+1):
                starts = await asyncio.gather(*[self.request('POST', f'/quizzes/{iid}/start') for _ in range(3)])
                self.assertTrue(all(r.status_code == 200 for r in starts), [r.text for r in starts])
                self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM quiz_starts WHERE quiz_id=$1', iid), number)
                refreshed = (await self.request('GET', f'/quizzes/{iid}')).json()
                self.assertEqual(refreshed['attempt_number'], number)
                self.assertTrue(all('is_correct' not in o for q in refreshed['questions'] for o in q['options']))
                if number > 1:
                    stale = await self.request('PATCH', f'/quizzes/{iid}/answers', json={'answers': [], 'attempt_number': number-1})
                    self.assertEqual(stale.status_code, 409)
                self.assertEqual((await self.submit_correct(iid, number)).status_code, 200)
            self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 403)
            history = (await self.request('GET', f'/quizzes/{iid}')).json()['attempt_history']
            self.assertEqual(len(history), limit)
            old = (await self.request('GET', f'/quizzes/{iid}?attempt_number=1')).json()
            self.assertEqual(old['attempt']['id'], history[-1]['id'])
            grades = (await self.request('GET', '/grades/my')).json()[0]
            self.assertEqual(grades['completed_items'], limit)

    async def test_content_access_does_not_apply_topical_attempt_rules(self):
        group = (await self.request('POST', '/admin/groups', 'admin', json={
            'code': 'SCOPE', 'title': 'Access scope', 'direction': 'Test', 'stream': '1'})).json()
        await self.request('POST', f'/admin/courses/{self.course}/groups', 'admin', json={'group_id': group['id']})
        await self.request('POST', f'/admin/groups/{group["id"]}/students', 'admin', json={'student_id': self.users['student']['id']})
        lesson = await self.item('lesson')
        homework = await self.item('homework')
        due = dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=1)
        await self.pool.execute('UPDATE homework SET deadline_at=$1 WHERE item_id=$2', due, homework)
        await self.request('PUT', f'/courses/lessons/{lesson}', 'teacher', json={
            'content': 'Theory remains available', 'youtube_url': 'https://www.youtube.com/watch?v=scope-test'})
        for iid in (lesson, homework):
            payload = {'group_ids': [group['id']], 'student_ids': [self.users['student']['id']]}
            self.assertEqual((await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json=payload)).status_code, 200)
            # Old values may exist in the shared grant tables; do not erase or enforce them for content.
            for table in ('item_group_access', 'item_student_access'):
                await self.pool.execute(f"UPDATE {table} SET deadline_at=now()-interval '1 day' WHERE item_id=$1", iid)
            before = await self.pool.fetchrow('SELECT * FROM item_student_access WHERE item_id=$1', iid)
            self.assertEqual((await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json=payload)).status_code, 200)
            self.assertEqual(dict(await self.pool.fetchrow('SELECT * FROM item_student_access WHERE item_id=$1', iid)), dict(before))
            denied = await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json=payload | {'quiz_rules': {'max_attempts': 3}})
            self.assertEqual(denied.status_code, 422)
            # Both independent student grants and group-only grants remain usable.
            for grants in ({'group_ids': [], 'student_ids': [self.users['student']['id']]}, {'group_ids': [group['id']], 'student_ids': []}):
                self.assertEqual((await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json=grants)).status_code, 200)
                path = f'/courses/lessons/{iid}' if iid == lesson else f'/homework/{iid}'
                for _ in range(3):
                    response = await self.request('GET', path)
                    self.assertEqual(response.status_code, 200, response.text)
                    if iid == lesson:
                        self.assertEqual(response.json()['content'], 'Theory remains available')
                        self.assertEqual(response.json()['youtube_url'], 'https://www.youtube.com/watch?v=scope-test')
                    else:
                        self.assertEqual(dt.datetime.fromisoformat(response.json()['deadline_at']), due)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM quiz_starts WHERE quiz_id=$1', lesson), 0)
        topical = await self.new_quiz(2)
        detail = (await self.request('GET', f'/courses/{self.course}', 'teacher')).json()
        for item in detail['items']:
            self.assertEqual('quiz_access_rules' in item, item['id'] == topical)
        # A lesson's independent embedded quiz keeps its original rules, not topical grant settings.
        self.assertEqual((await self.request('PUT', f'/quizzes/{lesson}/settings', 'teacher', json={'max_score': 100, 'weight_pct': 0})).status_code, 200)
        await self.question(lesson)
        view = (await self.request('GET', f'/quizzes/{lesson}')).json()
        self.assertFalse(view['is_topical'])
        self.assertIsNone(view['max_attempts'])
        self.assertIsNone(view['access_deadline_at'])
        self.assertEqual((await self.request('POST', f'/quizzes/{lesson}/start')).status_code, 200)

    async def test_deadline_validation_timezone_and_frozen_active_timer(self):
        future = (dt.datetime.now(dt.timezone.utc)+dt.timedelta(hours=1)).isoformat()
        iid = await self.new_quiz(2, future)
        await self.question(iid)
        for rules in [{'max_attempts': 0}, {'max_attempts': 1.5}, {'max_attempts': 1, 'deadline_at': '2020-01-01T00:00:00Z'}, {'max_attempts': 1, 'deadline_at': '2099-01-01T00:00:00'}]:
            response = await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json={'student_ids': [self.users['student']['id']], 'quiz_rules': rules})
            self.assertEqual(response.status_code, 422, response.text)
        await self.request('POST', f'/quizzes/{iid}/start')
        before = await self.pool.fetchval('SELECT deadline_at FROM quiz_starts WHERE quiz_id=$1', iid)
        later = (dt.datetime.now(dt.timezone.utc)+dt.timedelta(hours=2)).isoformat()
        await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json={'student_ids': [self.users['student']['id']], 'quiz_rules': {'max_attempts': 2, 'deadline_at': later}})
        self.assertEqual(await self.pool.fetchval('SELECT deadline_at FROM quiz_starts WHERE quiz_id=$1', iid), before)
        # Expired assignment blocks a NEW attempt even with remaining budget.
        await self.submit_correct(iid)
        await self.pool.execute("UPDATE item_student_access SET deadline_at=now()-interval '1 second' WHERE item_id=$1", iid)
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 403)

    async def test_expired_active_attempt_grades_only_saved_answers(self):
        iid = await self.new_quiz(2, (dt.datetime.now(dt.timezone.utc)+dt.timedelta(minutes=5)).isoformat())
        await self.question(iid, 4)
        await self.request('POST', f'/quizzes/{iid}/start')
        owner = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()['questions'][0]
        answer = {'question_id': owner['id'], 'selected_option_ids': [o['id'] for o in owner['options'] if o['is_correct']]}
        saved = await self.request('PATCH', f'/quizzes/{iid}/answers', json={'answers': [answer], 'revision': 0, 'request_id': 'topical-save-request-00001'})
        self.assertEqual(saved.status_code, 200, saved.text)
        await self.pool.execute("UPDATE quiz_starts SET deadline_at=now()-interval '1 second' WHERE quiz_id=$1", iid)
        await finalize_expired_quizzes(self.pool)
        late = await self.request('POST', f'/quizzes/{iid}/attempt', json={'answers': [], 'revision': 1})
        self.assertEqual(late.json()['auto_score'], 4)
        self.assertTrue(late.json()['submitted'])

    async def test_group_and_individual_permissions_intersect_not_add(self):
        iid = await self.new_quiz(2)
        await self.question(iid)
        group = (await self.request('POST', '/admin/groups', 'admin', json={'code': 'TQ', 'title': 'Test group', 'direction': 'Test', 'stream': '1'})).json()
        self.assertEqual((await self.request('POST', f'/admin/courses/{self.course}/groups', 'admin', json={'group_id': group['id']})).status_code, 200)
        await self.request('POST', f'/admin/groups/{group["id"]}/students', 'admin', json={'student_id': self.users['student']['id']})
        await self.pool.execute('INSERT INTO item_group_access(item_id,group_id,max_attempts) VALUES($1,$2,3)', iid, group['id'])
        view = (await self.request('GET', f'/quizzes/{iid}')).json()
        self.assertEqual(view['max_attempts'], 2)
        for number in (1, 2):
            await self.request('POST', f'/quizzes/{iid}/start')
            self.assertEqual((await self.submit_correct(iid, number)).status_code, 200)
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 403)

    async def test_group_only_access_and_expired_group_deadline(self):
        iid = await self.new_quiz(3)
        await self.question(iid)
        group = (await self.request('POST', '/admin/groups', 'admin', json={'code': 'GR', 'title': 'Group only', 'direction': 'Test', 'stream': '1'})).json()
        await self.request('POST', f'/admin/courses/{self.course}/groups', 'admin', json={'group_id': group['id']})
        await self.request('POST', f'/admin/groups/{group["id"]}/students', 'admin', json={'student_id': self.users['student']['id']})
        deadline = (dt.datetime.now(dt.timezone.utc)+dt.timedelta(hours=1)).isoformat()
        response = await self.request('PUT', f'/courses/items/{iid}/access', 'teacher', json={'group_ids': [group['id']], 'quiz_rules': {'max_attempts': 3, 'deadline_at': deadline}})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(await self.pool.fetchval('SELECT count(*) FROM item_student_access WHERE item_id=$1', iid), 0)
        view = (await self.request('GET', f'/quizzes/{iid}')).json()
        self.assertEqual(view['max_attempts'], 3)
        await self.pool.execute("UPDATE item_group_access SET deadline_at=now()-interval '1 second' WHERE item_id=$1", iid)
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/start')).status_code, 403)

    async def test_multiple_correct_word_questions_grade_without_changing_keys(self):
        iid = await self.new_quiz()
        files = {'file': ('multiple.docx', document(['<question>Select two', '<variant_correct>One', '<variant_correct>Two', '<variant>Three']))}
        self.assertEqual((await self.request('POST', f'/quizzes/{iid}/import-docx?preview=false', 'teacher', files=files)).status_code, 200)
        owner = (await self.request('GET', f'/quizzes/{iid}', 'teacher')).json()
        self.assertEqual(owner['questions'][0]['type'], 'multiple')
        await self.request('POST', f'/quizzes/{iid}/start')
        self.assertEqual((await self.submit_correct(iid)).json()['auto_score'], 1)

    async def test_manual_points_and_best_retake_grade_keep_weight_and_history(self):
        iid = await self.new_quiz(2)
        await self.question(iid, 2)
        manual = await self.request('POST', f'/quizzes/{iid}/questions', 'teacher', json={'type': 'short_text', 'prompt': 'Explain', 'points': 3})
        self.assertEqual(manual.status_code, 200, manual.text)
        await self.pool.execute('UPDATE quizzes SET weight_pct=25 WHERE item_id=$1', iid)
        await self.request('POST', f'/quizzes/{iid}/start')
        first = await self.submit_correct(iid)
        self.assertEqual((first.json()['auto_score'], first.json()['status']), (2, 'pending_review'))
        answer = await self.pool.fetchval('SELECT id FROM attempt_answers WHERE attempt_id=$1 AND question_id=$2', first.json()['attempt_id'], manual.json()['id'])
        graded = await self.request('POST', f'/quizzes/answers/{answer}/grade', 'teacher', json={'awarded_points': 3})
        self.assertEqual(graded.status_code, 200, graded.text)
        original = dict(await self.pool.fetchrow('SELECT * FROM quiz_attempts WHERE id=$1', first.json()['attempt_id']))
        await self.request('POST', f'/quizzes/{iid}/start')
        await self.request('POST', f'/quizzes/{iid}/attempt', json={'attempt_number': 2, 'revision': 0, 'answers': []})
        self.assertEqual(dict(await self.pool.fetchrow('SELECT * FROM quiz_attempts WHERE id=$1', original['id'])), original)
        grades = (await self.request('GET', '/grades/my')).json()[0]
        self.assertEqual(grades['final_grade'], 25)
        self.assertEqual(grades['completed_items'], 1)

    async def test_upgrade_preserves_old_scores_active_timers_and_permissions(self):
        name = self.schema+'_old'
        await self.connection.execute(f'CREATE SCHEMA "{name}"')
        conn = await asyncpg.connect(DATABASE_URL, server_settings={'search_path': name})
        try:
            await conn.execute(SCHEMA)
            for migration in ENT_MIGRATIONS+RELEASE_MIGRATIONS:
                await conn.execute(migration)
            teacher = await conn.fetchval("INSERT INTO users(email,password_hash,full_name,role) VALUES('teacher@old.test','unused','Teacher','teacher') RETURNING id")
            student = await conn.fetchval("INSERT INTO users(email,password_hash,full_name,role) VALUES('student@old.test','unused','Student','student') RETURNING id")
            course = await conn.fetchval("INSERT INTO courses(teacher_id,title) VALUES($1,'Old course') RETURNING id", teacher)
            iid = await conn.fetchval("INSERT INTO course_items(course_id,type,title,position) VALUES($1,'quiz','Old',0) RETURNING id", course)
            await conn.execute('INSERT INTO quizzes(item_id,max_score,weight_pct) VALUES($1,80,25)', iid)
            await conn.execute("INSERT INTO quiz_starts(quiz_id,student_id,deadline_at) VALUES($1,$2,now()+interval '2 hours')", iid, student)
            await conn.execute('INSERT INTO item_student_access(item_id,student_id) VALUES($1,$2)', iid, student)
            await conn.execute('INSERT INTO quiz_attempts(quiz_id,student_id,auto_score,manual_score) VALUES($1,$2,37,5)', iid, student)
            before = dict(await conn.fetchrow('SELECT * FROM quiz_attempts'))
            timer = await conn.fetchval('SELECT deadline_at FROM quiz_starts')
            for migration in TOPICAL_MIGRATIONS:
                await conn.execute(migration)
            after = dict(await conn.fetchrow('SELECT * FROM quiz_attempts'))
            self.assertEqual({key: after[key] for key in before}, before)
            self.assertEqual(await conn.fetchval('SELECT deadline_at FROM quiz_starts'), timer)
            self.assertEqual(await conn.fetchval('SELECT scoring_mode FROM quizzes'), 'legacy')
            permission = await conn.fetchrow('SELECT max_attempts,deadline_at FROM item_student_access')
            self.assertEqual((permission['max_attempts'], permission['deadline_at']), (1, None))
        finally:
            await conn.close()
            await self.connection.execute(f'DROP SCHEMA "{name}" CASCADE')
