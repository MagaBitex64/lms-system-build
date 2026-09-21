import base64
import io
import unittest

from docx import Document

from routers.quizzes import parse_docx_questions, stable_shuffle


class QuizDocxTests(unittest.TestCase):
    def make_docx(self, lines: list[str]) -> bytes:
        document = Document()
        for line in lines:
            document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def test_parses_questions_and_correct_options(self):
        questions = parse_docx_questions(
            self.make_docx(
                [
                    "Тест атауы",
                    "<question>2 + 2 нешеге тең?",
                    "<variant_correct>4",
                    "<variant>3",
                    "<variant>5",
                    "<question>Қазақстанның астанасы?",
                    "<variant>Алматы",
                    "<variant_correct>Астана",
                ]
            )
        )

        self.assertEqual(len(questions), 2)
        self.assertTrue(questions[0]["options"][0]["is_correct"])
        self.assertTrue(questions[1]["options"][1]["is_correct"])

    def test_shuffle_is_stable_and_student_specific(self):
        values = [{"id": value} for value in range(1, 20)]
        first = stable_shuffle(values, "quiz:student-1")
        repeated = stable_shuffle(values, "quiz:student-1")
        another_student = stable_shuffle(values, "quiz:student-2")

        self.assertEqual(first, repeated)
        self.assertNotEqual(first, another_student)

    def test_image_after_question_is_attached_to_question(self):
        document = Document()
        document.add_paragraph("<question>Суретте не көрсетілген?")
        image = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        document.add_paragraph().add_run().add_picture(io.BytesIO(image))
        document.add_paragraph("<variant_correct>Дұрыс жауап")
        document.add_paragraph("<variant>Қате жауап")
        output = io.BytesIO()
        document.save(output)

        questions = parse_docx_questions(output.getvalue())

        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["image"]["content_type"], "image/png")
        self.assertGreater(len(questions[0]["image"]["content"]), 0)


if __name__ == "__main__":
    unittest.main()
