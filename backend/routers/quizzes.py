import datetime
import base64
import hashlib
import io
import os
import re
import json
from urllib.parse import quote

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from core.access import ensure_course_owner, ensure_item_access, get_item_or_404
from core.db import get_pool
from core.quiz_attempts import quiz_state, complete_quiz, normalized_answers, decoded, utcnow
from core.deps import get_current_user, require_student, require_teacher
from storage.local import FileValidationError, storage

router = APIRouter(prefix="/quizzes", tags=["quizzes"])


class QuizSettingsIn(BaseModel):
    max_score: int = Field(ge=1, le=1000)
    weight_pct: float = Field(ge=0, le=100)
    open_at: datetime.datetime | None = None
    deadline_at: datetime.datetime | None = None
    close_at: datetime.datetime | None = None
    time_limit_minutes: int | None = Field(default=None, ge=1, le=1440)
    shuffle_questions: bool = True
    shuffle_options: bool = True


class OptionIn(BaseModel):
    text: str = Field(default="", max_length=1000)
    image_file_id: int | None = None
    is_correct: bool = False


class QuestionIn(BaseModel):
    type: str = Field(pattern="^(single|multiple|short_text|long_text)$")
    prompt: str = Field(min_length=1, max_length=10000)
    image_file_id: int | None = None
    explanation: str = Field(default="", max_length=5000)
    points: int = Field(default=1, ge=1, le=100)
    options: list[OptionIn] = Field(default_factory=list, max_length=100)


class AnswerIn(BaseModel):
    question_id: int
    selected_option_ids: list[int] = Field(default_factory=list, max_length=100)
    text_answer: str = Field(default="", max_length=20000)


class SubmitAttemptIn(BaseModel):
    answers: list[AnswerIn] = Field(max_length=1000)
    revision: int = Field(default=0, ge=0)
    request_id: str | None = Field(default=None, min_length=16, max_length=100)


class GradeAnswerIn(BaseModel):
    awarded_points: float = Field(ge=0)


async def get_quiz_or_404(pool, item_id: int):
    quiz = await pool.fetchrow("SELECT * FROM quizzes WHERE item_id = $1", item_id)
    if quiz is None:
        raise HTTPException(status_code=404, detail="Quiz not found")
    return quiz


async def ensure_quiz_row(pool, item_id: int):
    return await pool.fetchrow(
        """
        INSERT INTO quizzes (item_id)
        VALUES ($1)
        ON CONFLICT (item_id) DO UPDATE SET item_id = EXCLUDED.item_id
        RETURNING *
        """,
        item_id,
    )


def item_can_have_quiz(item: dict) -> bool:
    return item["type"] in ("lesson", "quiz")


def parse_docx_questions(content: bytes) -> list[dict]:
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Word файлын оқу мүмкін болмады. Дұрыс .docx файлын жүктеңіз") from exc

    tag_pattern = re.compile(r"^\s*\\?<(?P<tag>question|variant_correct|variant)>\s*(?P<text>.*)$", re.IGNORECASE)
    questions: list[dict] = []
    current: dict | None = None
    last_target: tuple[str, int | None] | None = None

    def consume_line(raw_line: str) -> None:
        nonlocal current, last_target
        line = raw_line.strip()
        if not line:
            return
        match = tag_pattern.match(line)
        if match:
            tag = match.group("tag").lower()
            text = match.group("text").strip()
            if tag == "question":
                current = {"prompt": text, "options": [], "image": None}
                questions.append(current)
                last_target = ("prompt", None)
                return
            if current is None:
                return
            current["options"].append({"text": text, "is_correct": tag == "variant_correct"})
            last_target = ("option", len(current["options"]) - 1)
            return

        if current is not None and last_target is not None:
            target, index = last_target
            if target == "prompt":
                current["prompt"] = f'{current["prompt"]}\n{line}'.strip()
            elif index is not None:
                current["options"][index]["text"] = f'{current["options"][index]["text"]}\n{line}'.strip()

    def consume_paragraph(paragraph) -> None:
        nonlocal current
        for line in paragraph.text.splitlines():
            consume_line(line)
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed"))
            if not relationship_id or current is None:
                continue
            image_part = document.part.related_parts.get(relationship_id)
            if image_part is None:
                continue
            if current["image"] is not None:
                raise HTTPException(
                    status_code=422,
                    detail=f'№{len(questions)}: бір сұраққа әзірше тек бір сурет импорттауға болады',
                )
            filename = os.path.basename(str(image_part.partname)) or f"question-{len(questions)}.png"
            current["image"] = {
                "filename": filename,
                "content_type": image_part.content_type,
                "content": image_part.blob,
            }

    for paragraph in document.paragraphs:
        consume_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    consume_paragraph(paragraph)

    if not questions:
        raise HTTPException(status_code=422, detail="Файлда <question> белгісі бар сұрақтар табылмады")
    errors = []
    for index, question in enumerate(questions, start=1):
        if not question["prompt"].strip():
            errors.append(f"№{index}: сұрақ мәтіні бос")
        if len(question["options"]) < 2:
            errors.append(f"№{index}: кемінде 2 жауап нұсқасы болуы керек")
        correct_count = sum(1 for option in question["options"] if option["is_correct"])
        if correct_count != 1:
            errors.append(f"№{index}: бір <variant_correct> болуы керек")
        if any(not option["text"].strip() for option in question["options"]):
            errors.append(f"№{index}: бос жауап нұсқасы бар")
    if errors:
        raise HTTPException(status_code=422, detail="; ".join(errors[:10]))
    return questions


def stable_shuffle(values: list[dict], seed: str) -> list[dict]:
    return sorted(values, key=lambda value: hashlib.sha256(f'{seed}:{value["id"]}'.encode()).digest())


# ---------- Teacher: quiz building ----------

@router.put("/{item_id}/settings")
async def update_quiz_settings(item_id: int, data: QuizSettingsIn, user: dict = Depends(require_teacher)):
    item = await get_item_or_404(item_id)
    await ensure_course_owner(user, item["course_id"])
    if not item_can_have_quiz(item):
        raise HTTPException(status_code=404, detail="This item cannot have a quiz")
    pool = await get_pool()
    await ensure_quiz_row(pool, item_id)
    await pool.execute(
        """
        UPDATE quizzes
        SET max_score = $1, weight_pct = $2, open_at = $3, deadline_at = $4,
            close_at = $5, time_limit_minutes = $6, shuffle_questions = $7,
            shuffle_options = $8
        WHERE item_id = $9
        """,
        data.max_score,
        data.weight_pct,
        data.open_at,
        data.deadline_at,
        data.close_at,
        data.time_limit_minutes,
        data.shuffle_questions,
        data.shuffle_options,
        item_id,
    )
    return {"ok": True, "message": "Quiz settings updated"}


@router.post("/{item_id}/questions")
async def add_question(item_id: int, data: QuestionIn, user: dict = Depends(require_teacher)):
    item = await get_item_or_404(item_id)
    await ensure_course_owner(user, item["course_id"])
    if not item_can_have_quiz(item):
        raise HTTPException(status_code=404, detail="This item cannot have a quiz")
    if data.type in ("single", "multiple"):
        if len(data.options) < 2:
            raise HTTPException(status_code=422, detail="Choice questions need at least 2 options")
        correct = sum(1 for o in data.options if o.is_correct)
        if data.type == "single" and correct != 1:
            raise HTTPException(status_code=422, detail="Single-choice questions need exactly 1 correct option")
        if data.type == "multiple" and correct < 1:
            raise HTTPException(status_code=422, detail="Multiple-choice questions need at least 1 correct option")
        if any(not o.text.strip() and o.image_file_id is None for o in data.options):
            raise HTTPException(status_code=422, detail="Each option needs text or an image")
    pool = await get_pool()
    image_file_ids = {
        file_id
        for file_id in [data.image_file_id, *(option.image_file_id for option in data.options)]
        if file_id is not None
    }
    if image_file_ids:
        owned_image_count = await pool.fetchval(
            """
            SELECT COUNT(*) FROM files
            WHERE id = ANY($1::bigint[]) AND (owner_id = $2 OR $3) AND mime LIKE 'image/%'
            """,
            list(image_file_ids),
            user["id"],
            user["role"] == "admin",
        )
        if owned_image_count != len(image_file_ids):
            raise HTTPException(status_code=422, detail="Question images must belong to the current teacher")
    await ensure_quiz_row(pool, item_id)
    if await pool.fetchval("SELECT EXISTS(SELECT 1 FROM quiz_attempts WHERE quiz_id=$1)", item_id):
        raise HTTPException(
            status_code=409,
            detail="Бұл тестті оқушылар тапсырып қойған. Нәтижелерді сақтау үшін сұрақтарды өзгертуге болмайды",
        )
    async with pool.acquire() as conn:
        async with conn.transaction():
            pos = await conn.fetchval(
                "SELECT COALESCE(MAX(position), -1) + 1 FROM quiz_questions WHERE quiz_id = $1", item_id
            )
            q = await conn.fetchrow(
                """
                INSERT INTO quiz_questions (quiz_id, type, prompt, image_file_id, explanation, points, position)
                VALUES ($1,$2,$3,$4,$5,$6,$7) RETURNING id
                """,
                item_id,
                data.type,
                data.prompt,
                data.image_file_id,
                data.explanation,
                data.points,
                pos,
            )
            for i, opt in enumerate(data.options):
                await conn.execute(
                    "INSERT INTO question_options (question_id, text, image_file_id, is_correct, position) VALUES ($1,$2,$3,$4,$5)",
                    q["id"],
                    opt.text,
                    opt.image_file_id,
                    opt.is_correct,
                    i,
                )
    return {"id": q["id"], "message": "Question added"}


@router.post("/{item_id}/import-docx")
async def import_questions_docx(
    item_id: int,
    preview: bool = True,
    file: UploadFile = File(...),
    user: dict = Depends(require_teacher),
):
    item = await get_item_or_404(item_id)
    await ensure_course_owner(user, item["course_id"])
    if not item_can_have_quiz(item):
        raise HTTPException(status_code=404, detail="Бұл элементке тест қосуға болмайды")
    if not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(status_code=422, detail="Тек .docx Word файлын жүктеңіз")
    chunks: list[bytes] = []
    content_size = 0
    while chunk := await file.read(1024 * 1024):
        content_size += len(chunk)
        if content_size > 10 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Word файлы 10 МБ-тан аспауы керек")
        chunks.append(chunk)
    content = b"".join(chunks)
    try:
        storage.validate_content(file.filename or "questions.docx", content)
    except FileValidationError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    questions = parse_docx_questions(content)
    if preview:
        preview_questions = []
        for question in questions:
            image = question.get("image")
            preview_questions.append({
                "prompt": question["prompt"],
                "options": question["options"],
                "image_data_url": (
                    f'data:{image["content_type"]};base64,{base64.b64encode(image["content"]).decode()}'
                    if image else None
                ),
            })
        return {"questions": preview_questions, "count": len(questions)}

    pool = await get_pool()
    await ensure_quiz_row(pool, item_id)
    has_attempts = await pool.fetchval("SELECT EXISTS(SELECT 1 FROM quiz_attempts WHERE quiz_id=$1)", item_id)
    if has_attempts:
        raise HTTPException(
            status_code=409,
            detail="Бұл тестті оқушылар тапсырып қойған. Нәтижелерді сақтау үшін жаңа тест құрып, файлды соған импорттаңыз",
        )
    for index, question in enumerate(questions, start=1):
        image = question.get("image")
        if image:
            try:
                storage.validate(image["filename"], len(image["content"]))
            except FileValidationError as exc:
                raise HTTPException(status_code=422, detail=f"№{index}: сурет форматы қолдау таппайды ({exc})") from exc
    saved_files: list[str] = []
    async with pool.acquire() as conn:
        try:
            async with conn.transaction():
                from core.config import MAX_USER_FILES, MAX_USER_STORAGE_BYTES
                await conn.fetchval('SELECT id FROM users WHERE id=$1 FOR UPDATE', user['id'])
                quota = await conn.fetchrow('SELECT count(*) n,COALESCE(sum(size),0) bytes FROM files WHERE owner_id=$1', user['id'])
                images = [q['image'] for q in questions if q.get('image')]
                if quota['n']+len(images)>MAX_USER_FILES or int(quota['bytes'])+sum(len(i['content']) for i in images)>MAX_USER_STORAGE_BYTES:
                    raise HTTPException(413,'Upload storage quota exceeded')
                start_position = await conn.fetchval(
                    "SELECT COALESCE(MAX(position), -1) + 1 FROM quiz_questions WHERE quiz_id = $1", item_id
                )
                for offset, question in enumerate(questions):
                    image_file_id = None
                    image = question.get("image")
                    if image:
                        stored_name, mime = storage.save(image["filename"], image["content"])
                        saved_files.append(stored_name)
                        image_file_id = await conn.fetchval(
                            """INSERT INTO files (owner_id,original_name,stored_name,mime,size)
                               VALUES ($1,$2,$3,$4,$5) RETURNING id""",
                            user["id"], image["filename"], stored_name, mime, len(image["content"]),
                        )
                    question_id = await conn.fetchval(
                        """INSERT INTO quiz_questions (quiz_id,type,prompt,image_file_id,explanation,points,position)
                           VALUES ($1,'single',$2,$3,'',1,$4) RETURNING id""",
                        item_id,
                        question["prompt"],
                        image_file_id,
                        start_position + offset,
                    )
                    for position, option in enumerate(question["options"]):
                        await conn.execute(
                            """INSERT INTO question_options (question_id,text,is_correct,position)
                               VALUES ($1,$2,$3,$4)""",
                            question_id,
                            option["text"],
                            option["is_correct"],
                            position,
                        )
        except Exception:
            for stored_name in saved_files:
                storage.delete(stored_name)
            raise
    return {"ok": True, "imported": len(questions)}


@router.get("/{item_id}/export-docx")
async def export_questions_docx(item_id: int, user: dict = Depends(require_teacher)):
    item = await get_item_or_404(item_id)
    await ensure_course_owner(user, item["course_id"])
    pool = await get_pool()
    await get_quiz_or_404(pool, item_id)
    questions = await pool.fetch(
        "SELECT id,prompt,type,image_file_id FROM quiz_questions WHERE quiz_id=$1 ORDER BY position", item_id
    )

    document = DocxDocument()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(item["title"], level=0)
    document.add_paragraph(
        "Импорт форматы: <question>, <variant_correct>, <variant>. "
        "Дұрыс жауап <variant_correct> белгісімен көрсетілген."
    )
    for index, question in enumerate(questions, start=1):
        if index > 1:
            document.add_paragraph()
        document.add_paragraph(f'<question>{question["prompt"]}')
        if question["image_file_id"]:
            file_row = await pool.fetchrow("SELECT stored_name FROM files WHERE id=$1", question["image_file_id"])
            if file_row and storage.exists(file_row["stored_name"]):
                document.add_picture(storage.path_for(file_row["stored_name"]), width=Inches(5.8))
        options = await pool.fetch(
            "SELECT text,is_correct FROM question_options WHERE question_id=$1 ORDER BY position", question["id"]
        )
        for option in options:
            tag = "variant_correct" if option["is_correct"] else "variant"
            document.add_paragraph(f'<{tag}>{option["text"]}')

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    safe_name = quote(f'{item["title"]}.docx')
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=quiz.docx; filename*=UTF-8''{safe_name}"},
    )


@router.delete("/questions/{question_id}")
async def delete_question(question_id: int, user: dict = Depends(require_teacher)):
    pool = await get_pool()
    q = await pool.fetchrow("SELECT * FROM quiz_questions WHERE id = $1", question_id)
    if q is None:
        raise HTTPException(status_code=404, detail="Question not found")
    item = await get_item_or_404(q["quiz_id"])
    await ensure_course_owner(user, item["course_id"])
    
    if await pool.fetchval("SELECT EXISTS(SELECT 1 FROM quiz_attempts WHERE quiz_id=$1)", q["quiz_id"]):
        raise HTTPException(
            status_code=409,
            detail="Бұл тестті оқушылар тапсырып қойған. Нәтижелерді сақтау үшін сұрақтарды өшіруге болмайды",
        )
    await pool.execute("DELETE FROM quiz_questions WHERE id = $1", question_id)
    
    return {"ok": True, "message": "Question deleted"}


# ---------- Quiz detail ----------

@router.get("/{item_id}")
async def get_quiz(item_id: int, user: dict = Depends(get_current_user)):
    item, course = await ensure_item_access(user, item_id)
    if not item_can_have_quiz(item):
        raise HTTPException(status_code=404, detail="Not a quiz")
    pool = await get_pool()
    quiz = await get_quiz_or_404(pool, item_id)
    questions = await pool.fetch(
        "SELECT * FROM quiz_questions WHERE quiz_id = $1 ORDER BY position", item_id
    )
    is_owner = user["role"] == "admin" or (user["role"] == "teacher" and course["teacher_id"] == user["id"])

    attempt = None
    state = None
    if user['role'] == 'student':
        state = await quiz_state(pool, item_id, user['id'])
        attempt = await pool.fetchrow('SELECT * FROM quiz_attempts WHERE quiz_id=$1 AND student_id=$2',item_id,user['id'])
    started = is_owner or state is not None or attempt is not None
    question_count = len(questions)
    total_points = sum(q['points'] for q in questions)
    if not started:
        questions = []

    out_questions = []
    question_ids = [q["id"] for q in questions]
    option_rows = await pool.fetch(
        """
        SELECT id, question_id, text, image_file_id, is_correct, position
        FROM question_options
        WHERE question_id = ANY($1::bigint[])
        ORDER BY question_id, position
        """,
        question_ids,
    ) if question_ids else []
    options_by_question: dict[int, list] = {}
    for option in option_rows:
        options_by_question.setdefault(option["question_id"], []).append(option)
    for q in questions:
        opts = options_by_question.get(q["id"], [])
        option_values = [dict(o) for o in opts]
        if user["role"] == "student" and quiz["shuffle_options"]:
            option_values = stable_shuffle(option_values, f'{item_id}:{user["id"]}:options:{q["id"]}')
        qd = {
            "id": q["id"],
            "type": q["type"],
            "prompt": q["prompt"],
            "image_file_id": q["image_file_id"],
            "points": q["points"],
            "options": [
                {"id": o["id"], "text": o["text"], "image_file_id": o["image_file_id"]}
                | ({"is_correct": o["is_correct"]} if is_owner or attempt else {})
                for o in option_values
            ],
        }
        # Explanation only visible to owner or after submitting
        if is_owner or attempt:
            qd["explanation"] = q["explanation"]
        out_questions.append(qd)

    if user["role"] == "student" and quiz["shuffle_questions"]:
        out_questions = stable_shuffle(out_questions, f'{item_id}:{user["id"]}:questions')

    result = {
        "type": "quiz",
        "is_owner": is_owner,
        "course_subject": course["ent_subject"],
        "item": {"id": item["id"], "title": item["title"], "note": item["note"], "course_id": item["course_id"]},
        "max_score": quiz["max_score"],
        "weight_pct": float(quiz["weight_pct"]),
        "open_at": str(quiz["open_at"]) if quiz["open_at"] else None,
        "deadline_at": str(quiz["deadline_at"]) if quiz["deadline_at"] else None,
        "close_at": str(quiz["close_at"]) if quiz["close_at"] else None,
        "time_limit_minutes": quiz["time_limit_minutes"],
        "shuffle_questions": quiz["shuffle_questions"],
        "shuffle_options": quiz["shuffle_options"],
        "total_points": total_points,
        "question_count": question_count,
        "quiz_started": started,
        "server_time": utcnow().isoformat(),
        "attempt_deadline_at": str(state['deadline_at']) if state and state['deadline_at'] else None,
        "remaining_seconds": max(0,int((state['deadline_at']-utcnow()).total_seconds())) if state and state['deadline_at'] else None,
        "revision": state['response_revision'] if state else 0,
        "saved_answers": decoded(state['responses']) if state else {},
        "questions": out_questions,
        "attempt": None,
    }
    if attempt:
        answers = await pool.fetch("SELECT * FROM attempt_answers WHERE attempt_id = $1", attempt["id"])
        result["attempt"] = {
            "id": attempt["id"],
            "submitted_at": str(attempt["submitted_at"]),
            "auto_score": float(attempt["auto_score"]),
            "manual_score": float(attempt["manual_score"]) if attempt["manual_score"] is not None else None,
            "status": attempt["status"],
            "answers": [
                {
                    "question_id": a["question_id"],
                    "selected_option_ids": list(a["selected_option_ids"]),
                    "text_answer": a["text_answer"],
                    "awarded_points": float(a["awarded_points"]) if a["awarded_points"] is not None else None,
                }
                for a in answers
            ],
        }
    return result


# ---------- Student: attempt ----------

@router.post("/{item_id}/start")
async def start_quiz(item_id: int, user: dict = Depends(require_student)):
    item, _course = await ensure_item_access(user, item_id)
    if not item_can_have_quiz(item):
        raise HTTPException(404, 'Not a quiz')
    pool = await get_pool()
    async with pool.acquire() as conn, conn.transaction():
        quiz = await conn.fetchrow('SELECT * FROM quizzes WHERE item_id=$1 FOR UPDATE', item_id)
        if quiz is None:
            raise HTTPException(404,'Quiz not found')
        current = utcnow()
        if quiz['open_at'] and current < quiz['open_at']:
            raise HTTPException(403,'Quiz is not open yet')
        if quiz['close_at'] and current > quiz['close_at']:
            raise HTTPException(403,'Quiz is closed')
        if not await conn.fetchval('SELECT 1 FROM quiz_questions WHERE quiz_id=$1',item_id):
            raise HTTPException(409,'Quiz has no questions')
        deadline = current + datetime.timedelta(minutes=quiz['time_limit_minutes']) if quiz['time_limit_minutes'] else None
        if quiz['close_at']:
            deadline = min(deadline,quiz['close_at']) if deadline else quiz['close_at']
        await conn.execute('INSERT INTO quiz_starts(quiz_id,student_id,deadline_at) VALUES($1,$2,$3) ON CONFLICT DO NOTHING',item_id,user['id'],deadline)
    return {'ok': True}


async def save_quiz_answers(item_id, data, user, submit=False):
    await ensure_item_access(user, item_id)
    pool = await get_pool()
    async with pool.acquire() as conn, conn.transaction():
        start = await conn.fetchrow('SELECT * FROM quiz_starts WHERE quiz_id=$1 AND student_id=$2 FOR UPDATE',item_id,user['id'])
        if start is None:
            raise HTTPException(403,'Start the quiz first')
        existing = await conn.fetchrow('SELECT * FROM quiz_attempts WHERE quiz_id=$1 AND student_id=$2',item_id,user['id'])
        expired = bool(start['deadline_at'] and start['deadline_at']<=utcnow())
        if existing or expired:
            result = existing or await complete_quiz(conn,start)
        else:
            replay = data.request_id and data.request_id == start['last_save_request_id']
            if not replay and data.revision != start['response_revision']:
                raise HTTPException(409,'Answers changed in another request. Reload the server answers.')
            if not replay:
                responses = decoded(start['responses']) or {}
                responses.update(await normalized_answers(conn,item_id,data.answers))
                start = await conn.fetchrow('UPDATE quiz_starts SET responses=$3::jsonb,response_revision=response_revision+1,last_save_request_id=$4 WHERE quiz_id=$1 AND student_id=$2 RETURNING *',item_id,user['id'],json.dumps(responses),data.request_id)
            result = await complete_quiz(conn,start) if submit else None
        return {'ok': True,'revision':start['response_revision'],'submitted':result is not None,
                'attempt_id':result['id'] if result else None,'auto_score':float(result['auto_score']) if result else None,
                'status':result['status'] if result else 'in_progress','expired':expired}


@router.patch("/{item_id}/answers")
async def save_answers(item_id: int, data: SubmitAttemptIn, user: dict = Depends(require_student)):
    return await save_quiz_answers(item_id,data,user)


@router.post("/{item_id}/attempt")
async def submit_attempt(item_id: int, data: SubmitAttemptIn, user: dict = Depends(require_student)):
    return await save_quiz_answers(item_id,data,user,submit=True)


# ---------- Teacher: review manual answers ----------

@router.get("/{item_id}/attempts")
async def list_attempts(item_id: int, user: dict = Depends(require_teacher)):
    item = await get_item_or_404(item_id)
    await ensure_course_owner(user, item["course_id"])
    pool = await get_pool()
    rows = await pool.fetch(
        """
        SELECT qa.*, u.full_name AS student_name, u.email AS student_email, q.deadline_at
        FROM quiz_attempts qa JOIN users u ON u.id = qa.student_id
        JOIN quizzes q ON q.item_id = qa.quiz_id
        WHERE qa.quiz_id = $1 ORDER BY qa.submitted_at DESC
        """,
        item_id,
    )
    return [
        {
            "id": r["id"],
            "student_id": r["student_id"],
            "student_name": r["student_name"],
            "student_email": r["student_email"],
            "submitted_at": str(r["submitted_at"]),
            "auto_score": float(r["auto_score"]),
            "manual_score": float(r["manual_score"]) if r["manual_score"] is not None else None,
            "status": r["status"],
            "late": bool(r["deadline_at"] and r["submitted_at"] > r["deadline_at"]),
        }
        for r in rows
    ]


@router.get("/attempts/{attempt_id}")
async def attempt_detail(attempt_id: int, user: dict = Depends(require_teacher)):
    pool = await get_pool()
    attempt = await pool.fetchrow(
        """
        SELECT qa.*, u.full_name AS student_name FROM quiz_attempts qa
        JOIN users u ON u.id = qa.student_id WHERE qa.id = $1
        """,
        attempt_id,
    )
    if attempt is None:
        raise HTTPException(status_code=404, detail="Attempt not found")
    item = await get_item_or_404(attempt["quiz_id"])
    await ensure_course_owner(user, item["course_id"])
    course_subject = await pool.fetchval("SELECT ent_subject FROM courses WHERE id=$1", item["course_id"])
    answers = await pool.fetch(
        """
        SELECT aa.*, qq.type, qq.prompt, qq.points AS question_points
        FROM attempt_answers aa JOIN quiz_questions qq ON qq.id = aa.question_id
        WHERE aa.attempt_id = $1 ORDER BY qq.position
        """,
        attempt_id,
    )
    return {
        "id": attempt["id"],
        "student_name": attempt["student_name"],
        "quiz_title": item["title"],
        "course_subject": course_subject,
        "submitted_at": str(attempt["submitted_at"]),
        "auto_score": float(attempt["auto_score"]),
        "manual_score": float(attempt["manual_score"]) if attempt["manual_score"] is not None else None,
        "status": attempt["status"],
        "answers": [
            {
                "id": a["id"],
                "question_id": a["question_id"],
                "type": a["type"],
                "prompt": a["prompt"],
                "question_points": a["question_points"],
                "selected_option_ids": list(a["selected_option_ids"]),
                "text_answer": a["text_answer"],
                "awarded_points": float(a["awarded_points"]) if a["awarded_points"] is not None else None,
            }
            for a in answers
        ],
    }


@router.post("/answers/{answer_id}/grade")
async def grade_answer(answer_id: int, data: GradeAnswerIn, user: dict = Depends(require_teacher)):
    pool = await get_pool()
    ans = await pool.fetchrow(
        """
        SELECT aa.*, qa.quiz_id, qq.points AS max_points FROM attempt_answers aa
        JOIN quiz_attempts qa ON qa.id = aa.attempt_id
        JOIN quiz_questions qq ON qq.id = aa.question_id
        WHERE aa.id = $1
        """,
        answer_id,
    )
    if ans is None:
        raise HTTPException(status_code=404, detail="Answer not found")
    item = await get_item_or_404(ans["quiz_id"])
    await ensure_course_owner(user, item["course_id"])
    if data.awarded_points > ans["max_points"]:
        raise HTTPException(status_code=422, detail="Awarded points exceed the question maximum")

    async with pool.acquire() as conn:
        async with conn.transaction():
            await conn.fetchval("SELECT id FROM quiz_attempts WHERE id=$1 FOR UPDATE", ans["attempt_id"])
            await conn.execute(
                "UPDATE attempt_answers SET awarded_points = $1 WHERE id = $2", data.awarded_points, answer_id
            )
            # Recompute manual score; mark graded when all manual answers are graded
            pending = await conn.fetchval(
                """
                SELECT COUNT(*) FROM attempt_answers aa
                JOIN quiz_questions qq ON qq.id = aa.question_id
                WHERE aa.attempt_id = $1 AND qq.type IN ('short_text','long_text') AND aa.awarded_points IS NULL
                """,
                ans["attempt_id"],
            )
            manual_points = await conn.fetchval(
                """
                SELECT COALESCE(SUM(aa.awarded_points), 0) FROM attempt_answers aa
                JOIN quiz_questions qq ON qq.id = aa.question_id
                WHERE aa.attempt_id = $1 AND qq.type IN ('short_text','long_text')
                """,
                ans["attempt_id"],
            )
            quiz = await conn.fetchrow("SELECT * FROM quizzes WHERE item_id = $1", ans["quiz_id"])
            total_points = await conn.fetchval(
                "SELECT COALESCE(SUM(points), 1) FROM quiz_questions WHERE quiz_id = $1", ans["quiz_id"]
            )
            manual_score = round(float(manual_points) / float(total_points) * quiz["max_score"], 2)
            await conn.execute(
                "UPDATE quiz_attempts SET manual_score = $1, status = $2 WHERE id = $3",
                manual_score,
                "graded" if pending == 0 else "pending_review",
                ans["attempt_id"],
            )
    return {"ok": True, "pending_manual": pending}
